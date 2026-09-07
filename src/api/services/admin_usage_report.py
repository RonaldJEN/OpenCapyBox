"""管理员使用报表：从现有调用记录聚合，不改变运行时计量口径。"""

from datetime import date, datetime, time, timedelta, timezone
from io import BytesIO
from typing import Any

from fastapi import HTTPException
from sqlalchemy import case, func, or_
from sqlalchemy.orm import Session as DBSession

from src.api.models.auth_user import AuthUser
from src.api.models.llm_call_record import LLMCallRecord
from src.api.models.llm_model import LLMModel
from src.api.models.session import Session
from src.api.utils.timezone import get_timezone

BEIJING = timezone(timedelta(hours=8))
METRICS = ("calls", "error_calls", "input_tokens", "output_tokens", "total_tokens", "missing_usage_calls")
REPORT_NOTE = (
    "有调用即活跃；开放、活跃、零使用人数仅统计当前启用账号，用量包含停用账号。"
    "调用次数为已落库记录数，模型按会话当前配置归属。"
    "Token 为已记录用量；删除会话会影响历史统计。"
)


def resolve_period(start_date: date | None, end_date: date | None, as_of: datetime | None = None) -> dict:
    now = datetime.now(BEIJING)
    if (start_date is None) != (end_date is None):
        raise HTTPException(422, "请同时提供开始日期和截止日期")
    end_date = end_date or now.date()
    start_date = start_date or end_date - timedelta(days=6)
    if start_date > end_date:
        raise HTTPException(422, "开始日期不能晚于截止日期")
    if (end_date - start_date).days >= 366:
        raise HTTPException(422, "查询时间范围过长，请缩小范围（最多366天）")
    if end_date > now.date():
        raise HTTPException(422, "截止日期不能晚于今天")
    if as_of is not None and (as_of.tzinfo is None or as_of > now):
        raise HTTPException(422, "查询截止时间必须带时区且不能晚于当前时间")
    cutoff = (as_of or now).astimezone(BEIJING)
    return {"start_date": start_date.isoformat(), "end_date": end_date.isoformat(),
            "as_of": cutoff.isoformat(), "timezone": "Asia/Shanghai"}


def _storage_time(value: datetime) -> datetime:
    # The existing DateTime columns store naive configured-local time, not UTC.
    return value.astimezone(get_timezone()).replace(tzinfo=None)


def _display_time(value: datetime | None) -> str | None:
    return value.replace(tzinfo=get_timezone()).astimezone(BEIJING).isoformat() if value else None


def build_usage_report(db: DBSession, period: dict) -> dict[str, Any]:
    start = _storage_time(datetime.combine(date.fromisoformat(period["start_date"]), time.min, BEIJING))
    end = _storage_time(datetime.combine(date.fromisoformat(period["end_date"]) + timedelta(days=1), time.min, BEIJING))
    cutoff = _storage_time(datetime.fromisoformat(period["as_of"]))
    record = LLMCallRecord
    # Fetch only one row per account/model; request/response texts never enter the report.
    grouped = (
        db.query(
            Session.user_id, Session.model_id, LLMModel.display_name.label("model_name"),
            func.count(record.id).label("calls"),
            func.sum(case((record.response_error.isnot(None), 1), else_=0)).label("error_calls"),
            func.coalesce(func.sum(record.usage_prompt_tokens), 0).label("input_tokens"),
            func.coalesce(func.sum(record.usage_completion_tokens), 0).label("output_tokens"),
            func.coalesce(func.sum(record.usage_total_tokens), 0).label("total_tokens"),
            func.sum(case((or_(record.usage_prompt_tokens.is_(None), record.usage_completion_tokens.is_(None),
                                    record.usage_total_tokens.is_(None)), 1), else_=0)).label("missing_usage_calls"),
            func.min(record.created_at).label("first_call_at"),
            func.max(record.created_at).label("last_call_at"),
        )
        .select_from(record)
        .join(Session, Session.id == record.session_id)
        .outerjoin(LLMModel, LLMModel.model_id == Session.model_id)
        .filter(record.created_at >= start, record.created_at < end, record.created_at <= cutoff)
        .group_by(Session.user_id, Session.model_id, LLMModel.display_name).all()
    )
    users = db.query(AuthUser.user_id, AuthUser.username, AuthUser.enabled).all()
    roster = {u.user_id: u for u in users}
    accounts: dict[str, dict] = {}
    models: dict[str | None, dict] = {}
    details = []

    def account(user_id: str) -> dict:
        if user_id not in accounts:
            user = roster.get(user_id)
            accounts[user_id] = {
                "user_id": user_id, "account": user.username if user else user_id,
                "enabled": bool(user and user.enabled), "model_count": 0,
                **dict.fromkeys(METRICS, 0),
            }
        return accounts[user_id]

    for user in users:
        if user.enabled:
            account(user.user_id)
    for row in grouped:
        owner = account(row.user_id)
        model_name = row.model_name or row.model_id or "未配置模型"
        if row.model_id not in models:
            models[row.model_id] = {
                "model_id": row.model_id, "model": model_name, "active_accounts": 0,
                **dict.fromkeys(METRICS, 0),
            }
        model = models[row.model_id]
        owner["model_count"] += 1
        # grouped has exactly one row per user/model, so this is a period-distinct count.
        model["active_accounts"] += 1
        metrics = {key: int(getattr(row, key) or 0) for key in METRICS}
        for key in METRICS:
            owner[key] += metrics[key]
            model[key] += metrics[key]
        details.append({"user_id": owner["user_id"], "account": owner["account"], "enabled": owner["enabled"],
                        "model_id": row.model_id, "model": model_name, **metrics,
                        "first_call_at": _display_time(row.first_call_at), "last_call_at": _display_time(row.last_call_at)})

    summary = {key: sum(a[key] for a in accounts.values()) for key in METRICS}
    summary["open_accounts"] = sum(bool(u.enabled) for u in users)
    summary["active_accounts"] = sum(a["enabled"] and a["calls"] > 0 for a in accounts.values())
    summary["unused_accounts"] = summary["open_accounts"] - summary["active_accounts"]
    for owner in accounts.values():
        owner["usage_status"] = "有调用" if owner["calls"] else "零使用"
        if not owner["calls"]:
            details.append({"user_id": owner["user_id"], "account": owner["account"], "enabled": owner["enabled"],
                            "model_id": None, "model": "无使用记录", **dict.fromkeys(METRICS, 0),
                            "first_call_at": None, "last_call_at": None})
    for item in [*accounts.values(), *models.values()]:
        item["share"] = item["total_tokens"] / summary["total_tokens"] if summary["total_tokens"] else None
    return {"period": period, "note": REPORT_NOTE, "summary": summary,
            "accounts": list(accounts.values()), "models": list(models.values()), "details": details}


SORT_KEYS = {
    "accounts": {"account", "model_count", "share", "usage_status", *METRICS},
    "models": {"model", "active_accounts", "share", *METRICS},
    "details": {"account", "model", "first_call_at", "last_call_at", *METRICS},
}


def report_rows(report: dict, view: str, sort_by: str = "total_tokens", direction: str = "desc",
                account_filter: str = "", model_filter: str = "") -> list[dict]:
    if sort_by not in SORT_KEYS[view]:
        raise HTTPException(422, "不支持的排序字段")
    rows = report[view]
    if view == "details":
        rows = [row for row in rows
                if (not account_filter or account_filter.casefold() in (row["account"] + " " + row["user_id"]).casefold())
                and (not model_filter or model_filter.casefold() in (row["model"] + " " + (row["model_id"] or "")).casefold())]
    # Deterministic ties make paging stable, including equal/zero token amounts.
    rows = sorted(rows, key=lambda r: (r.get("user_id", ""), r.get("model_id") or ""))
    known = [row for row in rows if row.get(sort_by) is not None]
    unknown = [row for row in rows if row.get(sort_by) is None]
    return sorted(known, key=lambda row: row[sort_by], reverse=direction == "desc") + unknown


COMMON_COLUMNS = [("calls", "调用次数"), ("error_calls", "错误调用"), ("input_tokens", "输入 Token"),
                  ("output_tokens", "输出 Token"), ("total_tokens", "总 Token"), ("missing_usage_calls", "Token 数据不完整调用")]
EXPORT_COLUMNS = {
    "accounts": [("account", "账号"), ("user_id", "账号标识"), ("enabled", "账号状态"), ("model_count", "使用模型数"),
                 *COMMON_COLUMNS, ("share", "期间占比"), ("usage_status", "使用状态")],
    "models": [("model", "模型"), ("model_id", "模型标识"), ("active_accounts", "活跃账号"), *COMMON_COLUMNS, ("share", "期间占比")],
    "details": [("account", "账号"), ("user_id", "账号标识"), ("enabled", "账号状态"), ("model", "模型"), ("model_id", "模型标识"),
                *COMMON_COLUMNS, ("first_call_at", "期间首次调用时间（北京时间）"), ("last_call_at", "期间末次调用时间（北京时间）")],
}


def export_usage_workbook(report: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook(write_only=True)

    def append(sheet, values, *, header=False):
        cells = []
        for value in values:
            cell = WriteOnlyCell(sheet, value=value)
            if isinstance(value, str):
                cell.data_type = "s"  # User/model names are literal text, including names starting with '='.
            if header:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="536347")
            elif isinstance(value, float):
                cell.number_format = "0.00%"
            elif isinstance(value, int):
                cell.number_format = "#,##0"
            elif isinstance(value, datetime):
                cell.number_format = "yyyy-mm-dd hh:mm:ss"
            cells.append(cell)
        sheet.append(cells)

    period = report["period"]
    period_label = f'{period["start_date"]} 至 {period["end_date"]}'
    for view, title in (("accounts", "汇总"), ("models", "模型汇总"), ("details", "账号×模型明细")):
        sheet = workbook.create_sheet(title)
        columns = [("period", "期间"), *EXPORT_COLUMNS[view]]
        header_row = 8 if view == "accounts" else 5
        sheet.freeze_panes = f"D{header_row + 1}"
        for index, (key, _) in enumerate(columns, 1):
            sheet.column_dimensions[get_column_letter(index)].width = 28 if key in {"period", "account", "model", "first_call_at", "last_call_at"} else 22
        append(sheet, ["OpenCapyBox 使用数据报表", period_label])
        append(sheet, ["查询截止时间（北京时间）", datetime.fromisoformat(period["as_of"]).replace(tzinfo=None)])
        append(sheet, ["统计口径", report["note"]])
        append(sheet, ["Token 数据不完整调用", report["summary"]["missing_usage_calls"]])
        if view == "accounts":
            summary_columns = [("open_accounts", "当前开放账号"), ("active_accounts", "期间活跃账号"),
                               ("unused_accounts", "期间零使用账号"), *COMMON_COLUMNS]
            append(sheet, [label for _, label in summary_columns], header=True)
            append(sheet, [report["summary"][key] for key, _ in summary_columns])
            append(sheet, [])
        append(sheet, [label for _, label in columns], header=True)
        rows = report_rows(report, view)
        sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(len(columns))}{header_row + len(rows)}"
        for row in rows:
            values = []
            for key, _ in columns:
                value = period_label if key == "period" else row.get(key)
                if key == "enabled":
                    value = "启用" if value else "停用或账号已不存在"
                elif key in {"first_call_at", "last_call_at"} and value:
                    value = datetime.fromisoformat(value).replace(tzinfo=None)
                values.append(value)
            append(sheet, values)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
