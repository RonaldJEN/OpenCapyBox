"""Critical report contracts on PostgreSQL, using an isolated rolled-back schema."""

import os
from datetime import date, datetime, timedelta
from io import BytesIO
from pathlib import Path
from uuid import uuid4

import pytest
from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine, text
from sqlalchemy.orm import Session as DBSession

from src.api.deps import get_current_admin_user
from src.api.models.auth_user import AuthUser
from src.api.models.database import Base, get_db
from src.api.models.llm_call_record import LLMCallRecord
from src.api.models.llm_model import LLMModel
from src.api.models.round import Round
from src.api.models.session import Session
from src.api.routes.admin_usage import router
from src.api.services.admin_usage_report import (
    BEIJING, build_usage_report, export_usage_workbook, report_rows, resolve_period,
)
from tests.db_safety import ensure_safe_test_database_url, load_dotenv_database_url


@pytest.fixture
def report_db():
    url = os.environ["TEST_DATABASE_URL"]
    ensure_safe_test_database_url(url, load_dotenv_database_url(Path(__file__).parent.parent))
    engine = create_engine(url)
    with engine.connect() as connection:
        transaction = connection.begin()
        schema = "test_usage_" + uuid4().hex
        connection.execute(text(f'CREATE SCHEMA "{schema}"'))
        connection.execute(text(f'SET LOCAL search_path TO "{schema}"'))
        Base.metadata.create_all(connection, tables=[model.__table__ for model in (AuthUser, Session, Round, LLMModel, LLMCallRecord)])
        db = DBSession(bind=connection)
        try:
            yield db
        finally:
            db.close()
            transaction.rollback()
    engine.dispose()


def seed(db):
    db.add_all([AuthUser(user_id=uid, username=name, enabled=enabled) for uid, name, enabled in (
        ("a", "=1+1", True), ("zero", "无调用", True), ("disabled", "停用用户", False),
    )])
    db.add(LLMModel(model_id="m", display_name="模型M", provider="openai", api_base="http://test", api_key="test", model_name="m"))
    db.add_all([Session(id=uid, user_id=uid, model_id="m") for uid in ("a", "disabled")])
    db.flush()
    for index, (uid, when, tokens, error) in enumerate([
        ("a", datetime(2026, 1, 1), 11, "failed"),
        ("a", datetime(2026, 1, 2, 23, 59, 59), None, "failed"),
        ("disabled", datetime(2026, 1, 2, 12), 19, None),
        ("a", datetime(2026, 1, 3), 999, None),
    ]):
        rid = str(index)
        db.add(Round(id=rid, session_id=uid, thread_id=uid, user_message="test"))
        db.flush()
        db.add(LLMCallRecord(session_id=uid, round_id=rid, step_index=1, request_messages="[]", request_tools="[]",
                             created_at=when, usage_prompt_tokens=tokens, usage_completion_tokens=0,
                             usage_total_tokens=tokens, response_error=error))
    db.flush()


def test_period_accounting_empty_accounts_and_excel(report_db):
    seed(report_db)
    report = build_usage_report(report_db, resolve_period(date(2026, 1, 1), date(2026, 1, 2)))
    assert report["summary"] == dict(calls=3, error_calls=2, input_tokens=30, output_tokens=0,
                                   total_tokens=30, missing_usage_calls=1, open_accounts=2, active_accounts=1, unused_accounts=1)
    assert report["models"][0]["active_accounts"] == 2  # a across two days still counts once
    assert sum(row["total_tokens"] for row in report["accounts"]) == sum(row["total_tokens"] for row in report["details"]) == 30
    zero = next(row for row in report["details"] if row["user_id"] == "zero")
    assert zero["calls"] == 0 and zero["model"] == "无使用记录" and zero["first_call_at"] is None
    assert [row["user_id"] for row in report_rows(report, "details", account_filter="停用")] == ["disabled"]
    workbook = load_workbook(BytesIO(export_usage_workbook(report)))
    assert workbook.sheetnames == ["汇总", "模型汇总", "账号×模型明细"]
    account_rows = list(workbook["汇总"].iter_rows(min_row=9))
    name_cell = next(row[1] for row in account_rows if row[2].value == "a")
    assert name_cell.value == "=1+1" and name_cell.data_type == "s"
    assert sum(row[9].value for row in account_rows) == 30
    assert workbook["汇总"].freeze_panes == "D9"
    assert workbook["模型汇总"]["K6"].number_format == "0.00%"
    assert isinstance(workbook["账号×模型明细"]["M6"].value, datetime)


def test_no_calls_keeps_open_accounts_and_cutoff_excludes_later_calls(report_db):
    seed(report_db)
    report = build_usage_report(report_db, resolve_period(date(2025, 12, 1), date(2025, 12, 2)))
    assert report["summary"]["open_accounts"] == report["summary"]["unused_accounts"] == 2
    assert len(report["accounts"]) == len(report["details"]) == 2
    assert report["models"] == [] and all(row["share"] is None for row in report["accounts"])
    cutoff = datetime(2026, 1, 1, 12, tzinfo=BEIJING)
    report = build_usage_report(report_db, resolve_period(date(2026, 1, 1), date(2026, 1, 2), cutoff))
    assert report["summary"]["calls"] == 1 and report["summary"]["total_tokens"] == 11


def test_report_routes_page_export_and_non_admin_permission(report_db):
    seed(report_db)
    app = FastAPI()
    app.include_router(router, prefix="/admin")
    app.dependency_overrides[get_db] = lambda: report_db
    app.dependency_overrides[get_current_admin_user] = lambda: "admin"
    params = {"start_date": "2026-01-01", "end_date": "2026-01-02", "page_size": 1}
    with TestClient(app) as client:
        response = client.get("/admin/usage-report", params=params)
        assert response.status_code == 200
        assert response.json()["total"] == 3 and len(response.json()["rows"]) == 1
        export = client.get("/admin/usage-report/export", params=params)
        assert export.status_code == 200
        assert len(list(load_workbook(BytesIO(export.content))["汇总"].iter_rows(min_row=9))) == 3
        assert "filename*=UTF-8''" in export.headers["content-disposition"]
        assert client.get("/admin/usage-report", params={**params, "sort_by": "request_messages"}).status_code == 422

        def deny():
            raise HTTPException(403, "需要管理员权限")

        app.dependency_overrides[get_current_admin_user] = deny
        for path in ("/admin/usage-report", "/admin/usage-report/export"):
            assert client.get(path, params=params).status_code == 403


def test_period_validation():
    assert resolve_period(date(2024, 1, 1), date(2024, 12, 31))["end_date"] == "2024-12-31"
    invalid = [(date(2024, 1, 1), date(2025, 1, 1)), (date(2024, 2, 1), date(2024, 1, 1)),
               (date(2024, 1, 1), None), (date.today(), date.today() + timedelta(days=2))]
    for start, end in invalid:
        with pytest.raises(HTTPException):
            resolve_period(start, end)
